import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt, RGBColor


def _plain(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, str):
        return value
    text = getattr(value, 'text', None)
    if text is not None and text is not value:
        return _plain(text)
    return str(value)


def _repair(value: str) -> str:
    """Repair common UTF-8/cp1252/latin1 mojibake without changing normal text."""
    s = _plain(value)
    markers = ('Ã', 'Â', 'â', 'ð', 'Ð', 'Ñ', '�')
    if not any(m in s for m in markers):
        return s

    best = s
    best_bad = sum(s.count(m) for m in markers)
    for enc in ('cp1252', 'latin1'):
        try:
            candidate = s.encode(enc).decode('utf-8')
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
        bad = sum(candidate.count(m) for m in markers)
        if bad < best_bad:
            best, best_bad = candidate, bad
    return best


def text(value: Any) -> str:
    s = _repair(value)
    s = s.replace('\u00a0', ' ')
    s = re.sub(r'[ \t]+', ' ', s)
    s = re.sub(r'\r\n?', '\n', s)
    return s.strip()


def safe(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', ' ', text(s)).strip()[:90] or 'quiz_test'


def _options(q):
    raw = q.get('options_json')
    if isinstance(raw, str):
        try:
            return json.loads(raw)
        except (TypeError, ValueError):
            return []
    return q.get('options') or []


def _correct_index(q) -> int:
    try:
        return int(q.get('correct_index'))
    except (TypeError, ValueError):
        return -1


def _add_option(doc: Document, letter: str, value: str, correct: bool) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f'{letter}) {text(value)}')
    if correct:
        run.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
        p.add_run('  ✓').bold = True


def export(quizzes, n=20, name='quiz_test'):
    """Export collected Telegram quiz records to DOCX files.

    Correct answers are taken only from the stored Telegram poll result
    (correct_index); this function never calls AI or changes the answer.
    """
    out = []
    if not quizzes:
        return out

    n = max(1, int(n))
    for k in range(0, len(quizzes), n):
        batch = quizzes[k:k + n]
        base = (
            name.strip()
            if name and name.strip() and name.strip().lower() != 'quiz_test'
            else text(batch[0].get('question', 'quiz_test'))[:70]
        )

        doc = Document()
        section = doc.sections[0]
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

        title = doc.add_paragraph()
        title.alignment = 1
        r = title.add_run(text(base))
        r.bold = True
        r.font.size = Pt(14)

        for i, q in enumerate(batch, 1):
            question = text(q.get('question', ''))
            qp = doc.add_paragraph()
            qp.paragraph_format.space_before = Pt(8)
            qp.paragraph_format.space_after = Pt(4)
            qr = qp.add_run(f'{i}. {question}')
            qr.bold = True
            qr.font.size = Pt(11)

            options = _options(q)
            correct = _correct_index(q)
            for j, option in enumerate(options):
                if j >= 26:
                    break
                _add_option(doc, chr(65 + j), option, j == correct)

            explanation = text(q.get('explanation', ''))
            if explanation:
                ep = doc.add_paragraph()
                ep.paragraph_format.space_before = Pt(2)
                er = ep.add_run('Izoh: ' + explanation)
                er.italic = True

            if i != len(batch):
                doc.add_paragraph().add_run().add_break(WD_BREAK.LINE)

        p = Path('exports') / f'{safe(base)}_{k // n + 1:02d}.docx'
        p.parent.mkdir(parents=True, exist_ok=True)
        doc.save(p)
        out.append(p)

    return out
